"""
VXR System Flow Reviewer — DOCX Generator

Builds a comprehensive study-guide DOCX explaining how the ViewXRent
(VXR) mobile app's tenant and landlord systems interact, and what
happens on the Supabase backend behind every screen.

Usage (from project root):
    python scripts\\generate_reviewer_docx.py

Output: VXR_System_Flow_Reviewer.docx in the project root.
"""

from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------

BRAND_ORANGE = RGBColor(0xFF, 0x70, 0x43)
BRAND_DARK = RGBColor(0x33, 0x33, 0x33)
BRAND_MUTED = RGBColor(0x66, 0x66, 0x66)
TABLE_HEADER_FILL = "FFE0B2"


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_page_break(doc) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        if level == 0:
            run.font.size = Pt(28)
            run.font.color.rgb = BRAND_ORANGE
        elif level == 1:
            run.font.size = Pt(20)
            run.font.color.rgb = BRAND_ORANGE
        elif level == 2:
            run.font.size = Pt(15)
            run.font.color.rgb = BRAND_DARK
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = BRAND_DARK


def add_paragraph(doc, text: str, bold: bool = False, italic: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_bullet(doc, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + level * 0.6)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_kv_bullet(doc, key: str, value: str, level: int = 0) -> None:
    """Bullet with a bolded leading term, e.g. **What:** description."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + level * 0.6)
    bold_run = p.add_run(f"{key} ")
    bold_run.bold = True
    bold_run.font.name = "Calibri"
    bold_run.font.size = Pt(11)
    text_run = p.add_run(value)
    text_run.font.name = "Calibri"
    text_run.font.size = Pt(11)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 6"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        set_cell_shading(hdr_cells[i], TABLE_HEADER_FILL)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    # spacing after the table
    doc.add_paragraph()
    return table


def add_toc(doc) -> None:
    """Insert a Word TOC field. The user presses F9 in Word to populate it."""
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Right-click and choose \"Update Field\" (or press F9) to refresh the Table of Contents."
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_sep)
    run._r.append(fld_text)
    run._r.append(fld_char_end)


# ---------------------------------------------------------------------------
# Document content
# ---------------------------------------------------------------------------

def build_cover(doc) -> None:
    # Vertical spacing for visual centering
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ViewXRent")
    run.font.name = "Calibri"
    run.font.size = Pt(40)
    run.font.color.rgb = BRAND_ORANGE
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Tenant & Landlord System Flow Reviewer")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = BRAND_DARK
    run.bold = True

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub2.add_run("How the Mobile App and Supabase Backend Work Together")
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = BRAND_MUTED
    run.italic = True

    for _ in range(4):
        doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Project: VXR (Flutter Mobile App)\n"
        "Backend: Supabase (Postgres + Auth + Storage + Edge Functions + Realtime)\n"
        "Payments: PayMongo (Philippines)\n"
        f"Document Date: {date.today().isoformat()}"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = BRAND_DARK

    add_page_break(doc)


def build_toc_page(doc) -> None:
    add_heading(doc, "Table of Contents", level=1)
    add_paragraph(
        doc,
        "This page contains a live Table of Contents. In Microsoft Word, right-click anywhere "
        "on the placeholder below and choose \"Update Field\" (or press F9) to populate it with "
        "page numbers from this document.",
        italic=True,
        size=10,
    )
    doc.add_paragraph()
    add_toc(doc)
    add_page_break(doc)


def build_part_1_overview(doc) -> None:
    add_heading(doc, "Part 1 — System Overview", level=1)

    add_heading(doc, "1.1  What is ViewXRent?", level=2)
    add_paragraph(
        doc,
        "ViewXRent (VXR) is a property rental platform that connects two kinds of people: "
        "tenants who are looking for a place to rent, and landlords who own properties they "
        "want to rent out. The whole experience happens inside a mobile app built with Flutter, "
        "which means the same code runs on both Android and iOS devices.",
    )
    add_paragraph(
        doc,
        "Instead of the app keeping all the data on the phone, everything important is stored "
        "on a backend service called Supabase. Supabase acts like the brain of the system: it "
        "stores user accounts, listings, applications, contracts, payments, and chat messages. "
        "It also enforces the rules about who is allowed to see what, so that tenants cannot "
        "peek at another tenant's data and landlords cannot mess with applications that "
        "belong to other landlords.",
    )
    add_paragraph(
        doc,
        "Payments are handled by a third service called PayMongo, which is a Philippine "
        "payment processor. VXR uses PayMongo for credit-card payments and for e-wallets "
        "like GCash, Maya, and GrabPay.",
    )

    add_heading(doc, "1.2  Tech Stack Summary", level=2)
    add_kv_bullet(doc, "Mobile app:", "Flutter (Dart). One codebase compiles to Android and iOS.")
    add_kv_bullet(doc, "Backend:", "Supabase. Provides the database (Postgres), authentication, file storage, server-side functions, and live realtime updates.")
    add_kv_bullet(doc, "Payments:", "PayMongo (sandbox / test mode in development). Used for both card and e-wallet payments. Edge Functions on Supabase talk to PayMongo using a secret key so the key never leaves the server.")
    add_kv_bullet(doc, "Maps:", "Google Maps Flutter plugin. Used on the search screen to show property pins, the user's location, and a 5 km radius circle.")
    add_kv_bullet(doc, "Location data:", "PSGC (Philippine Standard Geographic Code) JSON files. Powers the region / province / city / barangay pickers.")
    add_kv_bullet(doc, "Camera / 360°:", "camerawesome, camera_360, panorama_viewer, and opencv_dart for capturing and stitching 360° property tours on mobile.")
    add_kv_bullet(doc, "PDF / contracts:", "Flutter's pdf and printing packages for previewing and exporting signed lease contracts.")

    add_heading(doc, "1.3  Who Uses the App?", level=2)
    add_paragraph(
        doc,
        "There are three kinds of users in the system, but only two of them use the mobile "
        "app day-to-day:",
    )
    add_kv_bullet(doc, "Tenant:", "A person looking to rent a property. They browse listings, apply, sign a contract, pay rent, submit maintenance reports, and chat with the landlord.")
    add_kv_bullet(doc, "Landlord:", "A person who owns one or more properties and wants to rent them out. They list properties, review applications, sign contracts, receive payments, and manage their tenants.")
    add_kv_bullet(doc, "Admin:", "A platform operator who manages CMS content (FAQ pages, announcements, terms) and can step in to moderate. Admins work through admin-only tables and policies, not through the regular tenant/landlord UI.")
    add_paragraph(
        doc,
        "One important design choice: a single user account can be both a tenant and a "
        "landlord at the same time. The profile table tracks the user's role and a separate "
        "flag called is_landlord, which is flipped to true the first time the user creates a "
        "listing. This means someone can rent a place in one city while also renting out a "
        "property they own somewhere else, all from the same account.",
        italic=False,
    )

    add_heading(doc, "1.4  The Two-Sided Marketplace", level=2)
    add_paragraph(
        doc,
        "Think of VXR as a marketplace with three players: Tenant, Listing, and Landlord. "
        "The tenant and the landlord never interact directly through phone numbers or email "
        "addresses; everything goes through the app and the Supabase backend. The backend's "
        "job is to be the trust layer between them.",
    )
    add_paragraph(doc, "What the backend handles on behalf of both parties:", bold=True)
    add_bullet(doc, "Identity verification — confirms that the person on the other side is real (uploaded ID + selfie, AI-checked).")
    add_bullet(doc, "Listing integrity — only verified landlords can publish listings; admins can verify the listing documents.")
    add_bullet(doc, "Application gating — only verified tenants can apply for properties.")
    add_bullet(doc, "Contract recording — both digital signatures are timestamped and stored in the database.")
    add_bullet(doc, "Payment recording — every payment (online, offline, or mock) has a row in a tamper-evident ledger.")
    add_bullet(doc, "Communication — chat messages are stored on the server, so neither party can claim a message was never sent.")
    add_bullet(doc, "Access control — Row Level Security (RLS) policies make sure each user only sees their own data, no matter what query the app sends.")

    add_page_break(doc)


def build_part_2_tenant(doc) -> None:
    add_heading(doc, "Part 2 — Tenant Flow (End-to-End)", level=1)
    add_paragraph(
        doc,
        "This part walks through the tenant journey from creating an account all the way to "
        "moving out. Each subsection covers (a) what the tenant sees on screen, (b) what they "
        "do, and (c) what happens behind the scenes in the database and Supabase Edge "
        "Functions.",
    )

    # --- 2.1 Sign Up & Login ---
    add_heading(doc, "2.1  Sign Up and Login", level=2)
    add_paragraph(doc, "Screen: SignUpScreen, LoginScreen, LandingPage (splash).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Opens the app and sees a splash screen with two buttons: \"Get Started\" and \"Create Account.\"")
    add_bullet(doc, "On the sign-up screen, enters their full name, email, phone number, and password.")
    add_bullet(doc, "On the login screen, enters the email and password they registered with.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "The app calls Supabase Auth to create or sign in the user. Email and password go to the auth.users table managed by Supabase.")
    add_bullet(doc, "Full name and phone are saved into the user metadata of the auth.users row.")
    add_bullet(doc, "The app then runs ensureProfile(), which makes sure the matching row in the profiles table exists. role defaults to 'tenant' and is_landlord defaults to false.")
    add_bullet(doc, "Once signed in, the auth state stream tells the AuthGate to show the main HomeScreen instead of the login page.")
    add_paragraph(doc, "Result: a real, authenticated tenant session backed by a Supabase JWT.")

    # --- 2.2 Verification ---
    add_heading(doc, "2.2  Identity Verification", level=2)
    add_paragraph(doc, "Screen: VerificationScreen.", italic=True)
    add_paragraph(
        doc,
        "Before a tenant can apply for a listing, the system needs to know they are who they "
        "say they are. This is a one-time step and is enforced by a gate in the rental "
        "application flow.",
    )
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Picks the type of Philippine ID they will submit (passport, driver's license, national ID, etc.).")
    add_bullet(doc, "Takes a photo of the front of the ID using the rear camera.")
    add_bullet(doc, "If the ID has a back side, takes that photo too.")
    add_bullet(doc, "Takes a selfie using the front camera.")
    add_bullet(doc, "Reviews the photos and submits.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Each captured image is uploaded to a private Supabase Storage bucket (verifications).")
    add_bullet(doc, "The app calls a Supabase Edge Function called verify-identity. This function uses AI to compare the face on the ID with the selfie, read the name and document number using OCR, and assign a decision.")
    add_bullet(doc, "The decision is stored in the verifications table along with scores: ocr_confidence, face_match_score, name_match_score. Possible decisions are approved, rejected, or manual_review.")
    add_bullet(doc, "If approved, the profiles.is_verified flag is flipped to true. If it lands in manual_review, an admin reviews it later.")
    add_bullet(doc, "A row is also inserted into the notification table so the tenant sees an in-app notification with the result.")
    add_paragraph(doc, "Result: the tenant becomes eligible to apply for rental listings.")

    # --- 2.3 Browsing ---
    add_heading(doc, "2.3  Browsing Listings", level=2)
    add_paragraph(doc, "Screen: HomeScreen (home_page.dart), SearchScreen (search_field.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Scrolls through a feed of property cards on the home screen.")
    add_bullet(doc, "Switches to the search tab to filter results, type keywords, or view properties on a Google Map.")
    add_bullet(doc, "Taps a property to open the full details screen.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "fetchProperties() runs a SELECT on the listings table for all active listings, joined with their cover photo URL from listing_image.")
    add_bullet(doc, "The phone's geolocator gets the tenant's current location (with a fallback to Dasmariñas, Cavite if denied).")
    add_bullet(doc, "Distance from the tenant to each listing is calculated on the phone using the Haversine formula — no API call needed.")
    add_bullet(doc, "If the tenant is signed in, the app also fetches their bookmarks so that each card can show whether it's saved or not.")

    # --- 2.4 Filtering ---
    add_heading(doc, "2.4  Filtering and Searching", level=2)
    add_paragraph(doc, "Screen: FilterWidget (filter_widget.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Opens the filter drawer and adjusts sliders and toggles: number of bedrooms, bathrooms, area range, price range, furnishing level, pet policy, smoking policy.")
    add_bullet(doc, "Picks a sort order (relevance, lowest price, nearest).")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "All filtering and sorting happens on the phone. The list of listings is already loaded from Supabase; the filter just narrows down what is shown.")
    add_bullet(doc, "No backend calls happen during filtering. This makes the UI feel instant.")

    # --- 2.5 Bookmarking ---
    add_heading(doc, "2.5  Bookmarking (Favorites)", level=2)
    add_paragraph(doc, "Screen: FavoritesScreen (favorites_screen.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Taps the heart icon on a property card to save or unsave it.")
    add_bullet(doc, "Opens the Favorites tab to see all saved listings.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Tapping the heart calls toggleBookmark(). The app either INSERTs a row into the bookmarks table (user_id + listing_id) or DELETEs the existing row.")
    add_bullet(doc, "The Favorites screen queries the bookmarks table joined with listings to display the saved properties.")

    # --- 2.6 Unit Details ---
    add_heading(doc, "2.6  Viewing Unit Details", level=2)
    add_paragraph(doc, "Screen: UnitDetailsScreen (unit_details.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Reads the full property description, scrolls through photos, switches between Details, Amenities, and Location tabs.")
    add_bullet(doc, "Taps the 360° tour icon to view a panorama walkthrough.")
    add_bullet(doc, "Sees the landlord's profile name and verification badge.")
    add_bullet(doc, "Sees a CTA button that changes based on context: \"Apply for Rental\" if no application yet, \"Application Pending\" if one is being reviewed, or \"View Contract\" if approved.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "fetchListingDetails(listingId) loads the listing along with its pricing terms (listing_pricing) and lease terms (listing_lease_term).")
    add_bullet(doc, "fetchListingImages and fetchListingAmenities pull from listing_image and listing_amenities.")
    add_bullet(doc, "fetchPanoramaImages reads 360° images tagged with image_type='panorama' and room_label.")
    add_bullet(doc, "getMyApplicationForListing checks whether the current tenant already has an application row for this listing, and what its status is.")

    # --- 2.7 Application ---
    add_heading(doc, "2.7  Submitting a Rental Application", level=2)
    add_paragraph(doc, "Screen: RentalApplicationScreen (rental_application.dart), MyApplicationsScreen.", italic=True)
    add_paragraph(
        doc,
        "When a tenant taps \"Apply for Rental,\" the app runs a verification gate. If the "
        "tenant is not yet verified, they are routed to the verification screen first.",
    )
    add_paragraph(doc, "What the tenant does (5-step form):", bold=True)
    add_bullet(doc, "Step 1 — Personal Info: name, email, phone, date of birth, current address.")
    add_bullet(doc, "Step 2 — Employment: status (employed, self-employed, student, unemployed), company, job title, monthly income, length of employment, work address.")
    add_bullet(doc, "Step 3 — Rental History: first-time renter or not; if not, previous address, duration, reason for moving, previous landlord's name and contact.")
    add_bullet(doc, "Step 4 — Proof of Income: uploads a PDF or image of pay slips, contract, or bank statement.")
    add_bullet(doc, "Step 5 — Consent: ticks two checkboxes for identity verification consent and data privacy consent, then submits.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "submitRentalApplicationV2 INSERTs into the application table. status starts as 'pending'.")
    add_bullet(doc, "The income proof file is uploaded to the application-docs storage bucket; a corresponding row is inserted into application_documents linking the file URL to the application.")
    add_bullet(doc, "A database trigger creates a row in the notification table for the landlord, so they see the new application immediately.")

    # --- 2.8 Tracking Applications ---
    add_heading(doc, "2.8  Tracking My Applications", level=2)
    add_paragraph(doc, "Screen: MyApplicationsScreen (my_applications_screen.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Opens the My Applications screen from the profile menu.")
    add_bullet(doc, "Switches between Pending, Approved, and Rejected tabs.")
    add_bullet(doc, "If an application is still pending, can edit it or withdraw it.")
    add_bullet(doc, "If approved, sees a button that opens the contract for signing.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "fetchMyApplicationsWithListing() reads the application table filtered by tenant_id, joined with listings for context.")
    add_bullet(doc, "Editing updates the same row; withdrawing changes the status.")
    add_bullet(doc, "When the landlord later changes the status, a notification trigger fires and the tenant gets an in-app alert.")

    # --- 2.9 Contract ---
    add_heading(doc, "2.9  Reviewing and Signing the Lease Contract", level=2)
    add_paragraph(doc, "Screen: ContractViewScreen (contract_view_screen.dart).", italic=True)
    add_paragraph(
        doc,
        "After the landlord approves an application, the system creates a contract row "
        "automatically. The tenant can now open it from the My Applications screen.",
    )
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Reads the auto-filled lease contract (rent amount, deposit, lease duration, house rules).")
    add_bullet(doc, "Draws a signature on a signature pad and types in their printed name.")
    add_bullet(doc, "Taps \"Submit Signature.\"")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "signContractAsTenant updates the contract row with tenant_signature, tenant_printed_name, tenant_signed_at, and changes status from 'awaiting_tenant' to 'awaiting_landlord'.")
    add_bullet(doc, "A notification trigger alerts the landlord that it's their turn to sign.")
    add_bullet(doc, "Once the landlord signs, the contract status moves to 'fully_signed', and the tenant gets a notification that payment is now expected.")

    # --- 2.10 Move-in Payment ---
    add_heading(doc, "2.10  Move-in Payment (First-Time)", level=2)
    add_paragraph(doc, "Screen: ContractPaymentScreen (contract_payment_screen.dart).", italic=True)
    add_paragraph(
        doc,
        "The move-in payment is the first big payment in the lifecycle. It usually covers "
        "the first month's rent, the security deposit, and an advance month — depending on "
        "the listing's pricing terms.",
    )
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Reviews the amount due (rent + deposit + advance).")
    add_bullet(doc, "Picks a payment method: card, GCash, Maya, GrabPay, or bank transfer.")
    add_bullet(doc, "For cards: enters card details inside a PayMongo-managed form.")
    add_bullet(doc, "For e-wallets: gets redirected via a WebView to authorize the payment on the e-wallet provider's page.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "The app calls the paymongo-create-payment-intent Edge Function. The function calculates the move-in total from the contract's pricing fields, asks PayMongo to create a PaymentIntent, and returns the intent ID + client key + amount.")
    add_bullet(doc, "A pending row is seeded in the payment_transactions ledger.")
    add_bullet(doc, "The tenant attaches a payment method. For e-wallets, the user is redirected through PayMongo's hosted page; on return, the app calls paymongo-record-payment.")
    add_bullet(doc, "paymongo-record-payment fetches the latest state of the intent directly from PayMongo (trusting the source, not the browser). If status is 'succeeded', it upserts the payment table and updates payment_transactions to 'succeeded'.")
    add_bullet(doc, "For move-in (no billing_month), it also flips contract.status from 'fully_signed' to 'paid' and listings.status to 'rented'.")

    # --- 2.11 Recurring Rent ---
    add_heading(doc, "2.11  Recurring Monthly Rent Payment", level=2)
    add_paragraph(doc, "Screen: PaymentScreen (payment_screen.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Opens Payments tab and sees the next due month tile (e.g. \"Pay May 2026\").")
    add_bullet(doc, "Selects a saved payment method or enters a new one.")
    add_bullet(doc, "Pays the rent.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Same flow as move-in, but the Edge Function is called with billing_month set to a specific month. The amount is only the monthly_rent.")
    add_bullet(doc, "paymongo-record-payment marks the month as paid but does NOT change contract status, because the tenancy is already active.")
    add_bullet(doc, "A partial unique index in the payment table ensures only one 'succeeded' row per (contract, billing_month) — preventing double-payment.")

    # --- 2.12 In-Stay ---
    add_heading(doc, "2.12  The In-Stay Dashboard", level=2)
    add_paragraph(doc, "Screen: InStayDashboardScreen (in_stay_dashboard_screen.dart).", italic=True)
    add_paragraph(
        doc,
        "Once the move-in payment is successful, the tenant gains access to the In-Stay "
        "Dashboard. This is their command center for the duration of the tenancy.",
    )
    add_paragraph(doc, "What's on the dashboard:", bold=True)
    add_bullet(doc, "Active contract summary (property, start and end dates, rent amount).")
    add_bullet(doc, "Rent timeline showing each month's payment status (paid / unpaid / overdue).")
    add_bullet(doc, "Recent maintenance reports the tenant has filed.")
    add_bullet(doc, "Quick \"Contact Landlord\" button that opens the chat thread.")
    add_bullet(doc, "Option to request early termination.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "fetchMyActiveContract loads the contract row where status is 'paid', 'terminating', 'expiring', etc.")
    add_bullet(doc, "fetchRentMonths uses the contract's lease term and start date to compute the list of expected monthly payments. Each is cross-checked against the payment table to mark it paid or unpaid.")
    add_bullet(doc, "fetchMyReports lists the tenant's reports (open and in-progress).")
    add_bullet(doc, "getTermination checks contract_termination_request for any pending request.")

    # --- 2.13 Reports ---
    add_heading(doc, "2.13  Submitting Maintenance Reports", level=2)
    add_paragraph(doc, "Screen: ReportManagementScreen (report_management_screen.dart).", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Taps \"Submit Report\" from the dashboard.")
    add_bullet(doc, "Picks a type (maintenance, cleaning, noise, amenity, other) and a priority (low, medium, high).")
    add_bullet(doc, "Writes a title and description.")
    add_bullet(doc, "Optionally attaches photos showing the issue.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "submitReport INSERTs into the report table with status='open' and links to the contract, listing, tenant, and landlord.")
    add_bullet(doc, "Photo attachments are uploaded to a storage bucket and their URLs are saved on the report row.")
    add_bullet(doc, "A trigger creates a notification for the landlord.")

    # --- 2.14 Chat ---
    add_heading(doc, "2.14  Chatting with the Landlord", level=2)
    add_paragraph(doc, "Screen: ConversationsScreen, ChatThreadScreen.", italic=True)
    add_paragraph(doc, "What the tenant does:", bold=True)
    add_bullet(doc, "Opens the Messages tab to see all conversations they have with landlords.")
    add_bullet(doc, "Taps a conversation to open the thread.")
    add_bullet(doc, "Types and sends text messages, photos, or files (PDFs).")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Each conversation is a row in the conversation table that links a tenant and a landlord (sometimes scoped to a specific listing).")
    add_bullet(doc, "Each message goes into the message table with type='text' | 'image' | 'file'.")
    add_bullet(doc, "File attachments are uploaded to the conversation-attachments storage bucket.")
    add_bullet(doc, "Supabase Realtime broadcasts new rows in the message table to both participants instantly — so messages appear on the other phone without polling.")
    add_bullet(doc, "A database trigger keeps conversation.last_message and conversation.last_message_at up to date so the inbox shows the latest preview.")

    # --- 2.15 Termination ---
    add_heading(doc, "2.15  End-of-Tenancy", level=2)
    add_paragraph(
        doc,
        "A tenancy can end in several ways: the lease term expires, the tenant asks to move "
        "out early, or the landlord initiates termination.",
    )
    add_paragraph(doc, "What the tenant can do:", bold=True)
    add_bullet(doc, "From the In-Stay Dashboard, tap \"Request Early Termination.\"")
    add_bullet(doc, "Submit a reason; the request goes to the landlord for confirmation.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "A row is inserted into contract_termination_request.")
    add_bullet(doc, "Contract status moves from 'paid' to 'terminating'.")
    add_bullet(doc, "Once the landlord confirms vacancy and any final payments / refunds are settled, the contract moves to 'ended' or 'terminated'.")
    add_bullet(doc, "The listing can be re-listed and the tenant can apply for a different property.")

    add_page_break(doc)


def build_part_3_landlord(doc) -> None:
    add_heading(doc, "Part 3 — Landlord Flow (End-to-End)", level=1)
    add_paragraph(
        doc,
        "This part mirrors the tenant flow, but from the landlord's point of view. The same "
        "backend tables and Edge Functions appear, but the operations and access patterns are "
        "different — landlords write listings, approve applications, sign contracts, and "
        "record payments.",
    )

    add_heading(doc, "3.1  Becoming a Landlord", level=2)
    add_paragraph(doc, "What the user does:", bold=True)
    add_bullet(doc, "Signs up like any other tenant.")
    add_bullet(doc, "From the profile menu, taps \"Become a Landlord\" or \"Add a Listing.\"")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Before creating the first listing, the user is sent through identity verification (same flow as a tenant) — verified landlords keep the marketplace trustworthy.")
    add_bullet(doc, "When the first listing is saved, profiles.is_landlord is flipped to true.")
    add_bullet(doc, "The app's UI now exposes landlord-only screens like Manage Listings, Applicants, Contracts, and Tenant Management.")

    add_heading(doc, "3.2  Listing Enlistment", level=2)
    add_paragraph(doc, "Screen: HouseEnlistmentScreen, CreateListingScreen, ManageListingScreen.", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Reads the onboarding pitch about why to list on ViewXRent.")
    add_bullet(doc, "Fills out the listing form: title, description, address (using the PSGC picker), latitude / longitude (Google Maps pin), bedrooms, bathrooms, area, monthly rent, security deposit, advance, listing_type (lease or rent), amenities, pet and smoking policies.")
    add_bullet(doc, "Uploads ID and proof-of-ownership documents.")
    add_bullet(doc, "Saves the listing as a draft, then publishes.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "A new row is inserted into the listings table. Related rows go into listing_pricing, listing_lease_term, and listing_amenities.")
    add_bullet(doc, "Verification documents are stored in a private storage bucket. An admin reviews them and sets listings.is_verified='approved'.")
    add_bullet(doc, "Once approved, the listing appears in the public browse feed for tenants.")

    add_heading(doc, "3.3  Uploading Property Images", level=2)
    add_paragraph(doc, "Screen: ListingImageManagerScreen (listing_image_manager.dart).", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Picks images from the gallery or takes new photos with the camera.")
    add_bullet(doc, "Reorders thumbnails to control the carousel order.")
    add_bullet(doc, "Sets one image as the cover photo (the one shown on listing cards).")
    add_bullet(doc, "Removes unwanted images.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Image bytes are uploaded to the listing-images storage bucket.")
    add_bullet(doc, "Each uploaded image creates a row in listing_image with type='normal', sort_order, is_cover, and the storage URL.")
    add_bullet(doc, "Marking an image as cover updates listings.cover_photo_url to point to that file.")

    add_heading(doc, "3.4  Capturing 360° Panorama Tours", level=2)
    add_paragraph(doc, "Screen: PanoramaCaptureScreen, PanoramaManagerScreen. Mobile-only feature.", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Opens the panorama capture screen for a specific room.")
    add_bullet(doc, "Slowly rotates the phone left-to-right. The phone's gyroscope automatically triggers captures every ~10 degrees.")
    add_bullet(doc, "If they rotate too fast, the system skips frames to avoid motion blur.")
    add_bullet(doc, "Names each room (Living Room, Master Bedroom, Kitchen, etc.).")
    add_bullet(doc, "Adds, reorders, or deletes rooms as needed.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "Frames are stitched together on the device using OpenCV (opencv_dart). No cloud processing.")
    add_bullet(doc, "The final stitched image is uploaded to the listing-images storage bucket and inserted into listing_image with image_type='panorama' and a room_label.")
    add_bullet(doc, "Tenants can later open the 360° tour via the PanoramaTourViewer.")

    add_heading(doc, "3.5  Reviewing Applications", level=2)
    add_paragraph(doc, "Screen: ApplicantsScreen → ApplicationsScreen (enlistment_application.dart).", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Opens the Applicants tab and sees their listings with badges showing how many pending applications each has.")
    add_bullet(doc, "Taps a listing to see all applicants, filtered by status (Pending, Approved, Rejected).")
    add_bullet(doc, "Opens an application to review the tenant's personal info, employment, rental history, and proof of income.")
    add_bullet(doc, "Approves or rejects the application.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "fetchLandlordListings + fetchApplicationCounts power the badge counts.")
    add_bullet(doc, "Approving updates the application row's status to 'approved' and triggers getOrCreateContract, which creates a contract row in status 'awaiting_tenant'.")
    add_bullet(doc, "Rejecting just updates the status to 'rejected'.")
    add_bullet(doc, "In both cases, a notification trigger fires for the tenant.")

    add_heading(doc, "3.6  Signing the Lease Contract", level=2)
    add_paragraph(doc, "Screen: ContractViewScreen (shared with tenant).", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Waits for the tenant to sign first.")
    add_bullet(doc, "Opens the contract from the Contracts inbox.")
    add_bullet(doc, "Reviews the auto-filled terms and the tenant's signature.")
    add_bullet(doc, "Draws their own signature, types their printed name, and submits.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "signContractAsLandlord updates the contract row with landlord_signature, landlord_printed_name, landlord_signed_at, and changes status to 'fully_signed'.")
    add_bullet(doc, "The tenant gets a notification that the next step is move-in payment.")

    add_heading(doc, "3.7  Awaiting Move-in Payment", level=2)
    add_paragraph(doc, "Screen: LandlordContractsScreen (landlord_contracts_screen.dart).", italic=True)
    add_paragraph(
        doc,
        "Between contract signing and move-in payment, the landlord can monitor the contract "
        "from a dashboard.",
    )
    add_bullet(doc, "Contract appears in the \"Awaiting payment\" group.")
    add_bullet(doc, "Once the tenant successfully pays, the Edge Function paymongo-record-payment flips contract.status to 'paid' and listings.status to 'rented'.")
    add_bullet(doc, "The contract moves into the landlord's Tenant Management hub as an active tenant.")

    add_heading(doc, "3.8  The Tenant Management Hub", level=2)
    add_paragraph(doc, "Screen: TenantManagementScreen (tenantmanagement_screen.dart).", italic=True)
    add_paragraph(doc, "What the landlord sees:", bold=True)
    add_bullet(doc, "Hero banner with three stats: active tenants, open reports, properties.")
    add_bullet(doc, "Search bar that filters tenants by name, property, or email.")
    add_bullet(doc, "List of tenant cards showing avatar, property address, move-in date, days-in-stay, and open report count.")
    add_paragraph(doc, "What the landlord can do per tenant:", bold=True)
    add_bullet(doc, "View their payment history.")
    add_bullet(doc, "View their reports.")
    add_bullet(doc, "Open the chat thread.")
    add_bullet(doc, "Record an offline payment (cash, bank transfer, etc.).")
    add_bullet(doc, "Generate a PayMongo payment link to send to the tenant.")
    add_bullet(doc, "Initiate termination (mutual, non-renewal, eviction).")

    add_heading(doc, "3.9  Recording Offline Payments", level=2)
    add_paragraph(
        doc,
        "Sometimes the tenant pays by cash or direct bank transfer instead of going through "
        "the app. The landlord can log these payments so the ledger stays accurate.",
    )
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Opens the tenant card and taps \"Record payment.\"")
    add_bullet(doc, "Enters the amount in PHP.")
    add_bullet(doc, "Picks the method (cash, bank transfer, GCash, Maya, GrabPay, or other).")
    add_bullet(doc, "Optionally picks which billing month it covers.")
    add_bullet(doc, "Adds a note (receipt number, description).")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "The app calls the landlord-record-offline-payment Edge Function.")
    add_bullet(doc, "The function checks that the caller is the contract's landlord and that the contract is in an active state ('paid', 'terminating', 'expiring', 'terminated', 'ended').")
    add_bullet(doc, "It inserts rows into payment_transactions (primary ledger) and payment (legacy parity), with paymongo_payment_intent_id set to 'offline_<uuid>' so it's distinguishable from real PayMongo records.")
    add_bullet(doc, "The contract status is NOT changed — tenancy was already active before the offline payment.")

    add_heading(doc, "3.10  Generating Payment Links", level=2)
    add_paragraph(
        doc,
        "Instead of asking the tenant to open the app and pay rent, the landlord can send "
        "them a PayMongo hosted checkout URL by chat, email, or SMS.",
    )
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Opens \"Send link\" on a tenant card.")
    add_bullet(doc, "Picks the billing month and an optional note.")
    add_bullet(doc, "Taps \"Generate link.\"")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "The Edge Function paymongo-create-payment-link asks PayMongo to create a Link object with the rent amount.")
    add_bullet(doc, "A row is inserted into payment_links with the checkout URL, status 'pending', expires_at, and the linked contract / billing_month.")
    add_bullet(doc, "When the tenant pays the link, PayMongo fires the link.payment.paid webhook into paymongo-webhook, which marks the link as paid and writes payment + payment_transactions rows just like a normal in-app payment.")

    add_heading(doc, "3.11  Handling Reports", level=2)
    add_paragraph(doc, "Screen: ReportManagementScreen (landlord-side tabs).", italic=True)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Opens \"Incoming Reports\" and filters by type, status, or priority.")
    add_bullet(doc, "Reads the description and photos.")
    add_bullet(doc, "Marks the report as In Progress, then Resolved.")
    add_bullet(doc, "Optionally writes a response to the tenant.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "updateReportStatus changes report.status in the report table.")
    add_bullet(doc, "respondToReport saves landlord_response and landlord_responded_at.")
    add_bullet(doc, "Each change triggers a notification for the tenant.")

    add_heading(doc, "3.12  Chatting with Tenants", level=2)
    add_paragraph(
        doc,
        "Same conversation and message tables as the tenant side. The landlord sees their "
        "inbox of conversations; messages stream live via Supabase Realtime. File attachments "
        "are stored in the conversation-attachments bucket.",
    )

    add_heading(doc, "3.13  Initiating Termination", level=2)
    add_paragraph(doc, "What the landlord does:", bold=True)
    add_bullet(doc, "Picks the termination type: mutual (both sides agree), non-renewal (lease ending), or eviction (for cause).")
    add_bullet(doc, "Enters a reason (required).")
    add_bullet(doc, "Submits the request.")
    add_paragraph(doc, "What happens behind the scenes:", bold=True)
    add_bullet(doc, "requestTermination writes a row into contract_termination_request with the type, reason, effective_date, and security_deposit_amount.")
    add_bullet(doc, "Contract status moves to 'terminating'.")
    add_bullet(doc, "Tenant gets a notification and the In-Stay Dashboard shows the move-out checklist.")
    add_bullet(doc, "Once both parties confirm move-out, the contract moves to 'ended' or 'terminated'.")

    add_page_break(doc)


def build_part_4_backend(doc) -> None:
    add_heading(doc, "Part 4 — Backend Architecture", level=1)

    add_heading(doc, "4.1  Why Supabase?", level=2)
    add_paragraph(
        doc,
        "Supabase is an open-source alternative to Firebase. It bundles five things VXR needs "
        "into one product: a Postgres database, an authentication system, file storage, "
        "server-side functions (Edge Functions), and live realtime subscriptions on database "
        "tables. Using Supabase means VXR's small team does not have to run separate servers "
        "for each capability.",
    )

    add_heading(doc, "4.2  The Database — Module by Module", level=2)
    add_paragraph(
        doc,
        "VXR's database is split into modules, each represented by a SQL migration file in "
        "supabase/migrations/. The major tables and what they're for:",
    )
    add_heading(doc, "People tables", level=3)
    add_kv_bullet(doc, "profiles", "One row per user, mirroring auth.users. Holds full_name, phone, role (tenant/landlord/admin), is_landlord, is_verified, avatar_url, and verification document URLs.")
    add_kv_bullet(doc, "verifications", "Identity verification submissions. Stores image URLs (ID front, back, selfie), AI scores, and the decision (approved/rejected/manual_review).")
    add_heading(doc, "Property tables", level=3)
    add_kv_bullet(doc, "listings", "The core property table. Each row is one rentable unit owned by one landlord.")
    add_kv_bullet(doc, "listing_image", "All images attached to a listing — regular photos (type='normal') and 360° panoramas (image_type='panorama' with room_label).")
    add_kv_bullet(doc, "listing_pricing", "Pricing breakdown for a listing: monthly rent, security deposit, advance rent, association dues.")
    add_kv_bullet(doc, "listing_lease_term", "Lease duration rules: minimum stay, payment due date, grace period.")
    add_kv_bullet(doc, "listing_amenities", "Amenities checklist (wifi, parking, aircon, fully furnished, etc.).")
    add_heading(doc, "Engagement tables", level=3)
    add_kv_bullet(doc, "bookmarks", "Saved listings per tenant.")
    add_kv_bullet(doc, "application", "Rental applications, one row per (tenant, listing) attempt.")
    add_kv_bullet(doc, "application_documents", "Files attached to applications (proof of income, etc.).")
    add_heading(doc, "Contract tables", level=3)
    add_kv_bullet(doc, "contract", "One row per approved application. Stores signatures, lease terms, and the lifecycle status (awaiting_tenant → … → ended).")
    add_kv_bullet(doc, "contract_termination_request", "Termination requests with type (mutual / non-renewal / eviction), reason, effective date.")
    add_heading(doc, "Payment tables", level=3)
    add_kv_bullet(doc, "payment", "Legacy primary payment table (carried over from a prior Stripe-based implementation). Still written for compatibility with older queries.")
    add_kv_bullet(doc, "payment_methods", "Saved e-wallet shortcuts (GCash, Maya, GrabPay, bank transfer). Cards are NOT saved here because PayMongo PH does not expose card vaulting.")
    add_kv_bullet(doc, "payment_transactions", "The new per-attempt ledger. Tracks every payment attempt (succeeded, pending, failed, refunded). Supports offline payments via recorded_by.")
    add_kv_bullet(doc, "payment_links", "PayMongo hosted checkout links generated by landlords for tenants to pay specific rent months.")
    add_heading(doc, "Communication tables", level=3)
    add_kv_bullet(doc, "conversation", "One row per (tenant, landlord, optional listing) thread.")
    add_kv_bullet(doc, "message", "Individual messages. Types: text, image, file.")
    add_kv_bullet(doc, "notification", "Per-user notification queue. Filled by database triggers when important events happen.")
    add_heading(doc, "Maintenance & admin tables", level=3)
    add_kv_bullet(doc, "report", "Maintenance, cleaning, noise, and amenity reports filed by tenants.")
    add_kv_bullet(doc, "cms_page / cms_announcement / cms_faq", "Admin-managed static content: terms, privacy, help center, banners, FAQs.")

    add_heading(doc, "4.3  Row Level Security (RLS) Explained", level=2)
    add_paragraph(
        doc,
        "Every important table in the database has Row Level Security turned on. RLS is a "
        "Postgres feature where each row carries an ownership rule, and queries can only "
        "return rows that satisfy that rule for the current user's auth.uid().",
    )
    add_paragraph(doc, "Examples of policies enforced by RLS:", bold=True)
    add_bullet(doc, "Tenants can read their own bookmarks, but not anyone else's. SELECT on bookmarks is filtered to user_id = auth.uid().")
    add_bullet(doc, "Landlords can SELECT applications, but only those for listings where listings.landlord_id = auth.uid().")
    add_bullet(doc, "Tenants can SELECT a contract only if contracts.tenant_id = auth.uid(); landlords only if contracts.landlord_id = auth.uid().")
    add_bullet(doc, "Messages and conversations are visible only to the two participants.")
    add_bullet(doc, "Verification documents are owner-only — even other landlords cannot peek at a tenant's ID.")
    add_bullet(doc, "Admins use a helper function is_admin() that grants full read/write on certain tables (CMS, profiles, all listings).")
    add_paragraph(
        doc,
        "RLS is enforced at the database level. Even if the app accidentally sends a query "
        "asking for someone else's data, Postgres returns zero rows. This is the single most "
        "important safety feature of the system.",
    )

    add_heading(doc, "4.4  Storage Buckets", level=2)
    add_paragraph(
        doc,
        "Files are stored in Supabase Storage, which is essentially S3 with access policies "
        "tied to auth.uid().",
    )
    add_table(
        doc,
        headers=["Bucket", "Purpose", "Access"],
        rows=[
            ["property-images", "Landlord-uploaded photos of properties", "Public read for active listings; write requires landlord_id match"],
            ["listing-images", "Listing photos and panorama images", "Public read; write owner-only"],
            ["listing-contracts", "Generated and signed PDF contracts", "Parties only"],
            ["verifications", "Identity verification ID + selfie images", "Owner-only; admins via is_admin()"],
            ["panorama-tours", "Raw frames captured for stitching", "Owner-only"],
            ["application-docs", "Tenant uploads (proof of income)", "Tenant who uploaded + the listing's landlord"],
            ["conversation-attachments", "Files sent inside chat threads", "Conversation participants only"],
        ],
    )

    add_heading(doc, "4.5  Edge Functions Catalogue", level=2)
    add_paragraph(
        doc,
        "Edge Functions are small TypeScript programs that run on Supabase's serverless "
        "platform. They are useful for operations that should not be done by the mobile app "
        "directly — usually because they need a secret API key (like the PayMongo secret) or "
        "because they need to do multi-step writes that must succeed or fail as a unit.",
    )

    rows = [
        ["verify-identity", "Runs AI checks on uploaded ID + selfie, returns a decision, updates verifications and profiles.is_verified.", "VerificationScreen"],
        ["paymongo-create-payment-intent", "Creates a PayMongo PaymentIntent for either move-in (rent + deposit + advance) or recurring rent (single month). Seeds a pending row in payment_transactions.", "ContractPaymentScreen, PaymentScreen"],
        ["paymongo-attach-payment-method", "Attaches a saved or new payment method to an intent. Handles redirect flows for e-wallets.", "ContractPaymentScreen"],
        ["paymongo-record-payment", "Called after the payment redirect. Fetches the intent state directly from PayMongo, upserts payment and payment_transactions to 'succeeded', flips contract and listing status if it was a move-in.", "ContractPaymentScreen, PaymentScreen"],
        ["paymongo-record-mock-payment", "Skips PayMongo entirely. Writes a synthetic 'pi_mock_<uuid>' payment row. Enabled by MOCK_PAYMENTS_ENABLED for demos and tests.", "ContractPaymentScreen (mock mode)"],
        ["paymongo-add-method / delete-method / set-default-method", "CRUD on saved e-wallet shortcuts in the payment_methods table.", "PaymentScreen → Saved Methods"],
        ["paymongo-create-payment-link", "Generates a PayMongo Link checkout URL for a specific (contract, billing_month). Writes a row to payment_links.", "TenantManagementScreen → Send Link"],
        ["paymongo-webhook", "Public endpoint registered with PayMongo. Receives payment.paid, payment.failed, payment.refunded, link.payment.paid events. Idempotent: never downgrades a 'succeeded' row.", "PayMongo (server-to-server)"],
        ["landlord-record-offline-payment", "Records cash, bank, or direct-wallet payments outside PayMongo. Writes dual rows in payment and payment_transactions with paymongo_payment_intent_id='offline_<uuid>'.", "TenantManagementScreen → Record Payment"],
    ]
    add_table(doc, headers=["Edge Function", "What it does", "Called from"], rows=rows, col_widths=[4.5, 9.5, 4])

    add_heading(doc, "4.6  Realtime", level=2)
    add_paragraph(
        doc,
        "Supabase Realtime is a websocket-based system that streams INSERTs, UPDATEs, and "
        "DELETEs from the database to any connected client. VXR uses it for two things:",
    )
    add_bullet(doc, "Chat — both tenant and landlord subscribe to new messages on their conversations. Messages appear on the other phone within a fraction of a second, with no polling.")
    add_bullet(doc, "Notifications — the notification table is also live, so unread badges update instantly when a trigger inserts a new row.")

    add_heading(doc, "4.7  Notifications", level=2)
    add_paragraph(
        doc,
        "The notification table is filled by database triggers that fire automatically on "
        "important state changes. The app does not need to insert notifications manually.",
    )
    add_paragraph(doc, "Events that produce notifications:", bold=True)
    add_bullet(doc, "Message sent → notification for the recipient.")
    add_bullet(doc, "Application submitted → notification for the landlord.")
    add_bullet(doc, "Application approved / rejected → notification for the tenant.")
    add_bullet(doc, "Contract created / signed / cancelled → notifications for both parties.")
    add_bullet(doc, "Report submitted / status changed / landlord response → notifications.")
    add_bullet(doc, "Payment succeeded → notifications for tenant and landlord.")
    add_bullet(doc, "Verification decision → notification for the user.")
    add_bullet(doc, "Listing verification approved / rejected → notification for the landlord.")

    add_page_break(doc)


def build_part_5_payments(doc) -> None:
    add_heading(doc, "Part 5 — Payment System Deep Dive", level=1)

    add_heading(doc, "5.1  The PayMongo Model", level=2)
    add_paragraph(
        doc,
        "PayMongo is a Philippine payment processor. The flow it follows is similar to other "
        "modern processors like Stripe: the app first creates a PaymentIntent (a server-side "
        "object describing what is being charged), then the user picks a payment method, then "
        "PayMongo charges the method and reports the result back via a webhook.",
    )
    add_paragraph(doc, "VXR uses PayMongo in two modes:", bold=True)
    add_kv_bullet(doc, "Sandbox / test mode:", "Uses pk_test_… and sk_test_… keys; no real money moves. Used in development and demos.")
    add_kv_bullet(doc, "Production:", "Will use pk_live_… and sk_live_… keys once the platform launches.")

    add_heading(doc, "5.2  The Move-in Payment Lifecycle", level=2)
    add_paragraph(doc, "Step-by-step from contract signing to listing marked as rented:", bold=True)
    add_bullet(doc, "1. Both parties have signed; contract.status = 'fully_signed'.")
    add_bullet(doc, "2. Tenant opens ContractPaymentScreen.")
    add_bullet(doc, "3. App calls paymongo-create-payment-intent with the contract_id only (no billing_month).")
    add_bullet(doc, "4. Edge Function calculates the move-in total = monthly_rent + security_deposit + advance_rent. Asks PayMongo to create a PaymentIntent.")
    add_bullet(doc, "5. App receives intent ID + client key. A pending row is seeded in payment_transactions.")
    add_bullet(doc, "6. Tenant picks a method (card or e-wallet). For e-wallets, PayMongo redirects them through a WebView.")
    add_bullet(doc, "7. Once the redirect returns, the app calls paymongo-record-payment. Edge Function asks PayMongo for the current intent state.")
    add_bullet(doc, "8. If status is 'succeeded', the function upserts payment + payment_transactions, sets contract.status='paid', and sets listings.status='rented'.")
    add_bullet(doc, "9. Notification triggers fire for both parties.")
    add_bullet(doc, "10. Tenant is unlocked into the In-Stay Dashboard.")

    add_heading(doc, "5.3  The Monthly Rent Lifecycle", level=2)
    add_paragraph(doc, "How it differs from move-in:", bold=True)
    add_bullet(doc, "billing_month is set (e.g. '2026-05-01').")
    add_bullet(doc, "Amount is only the monthly rent, no deposit or advance.")
    add_bullet(doc, "Contract status does NOT change (tenancy is already active).")
    add_bullet(doc, "A partial unique index on payment ensures only one succeeded row per (contract, billing_month).")
    add_bullet(doc, "If the tenant tries to pay the same month twice, the Edge Function detects an existing succeeded row and refuses to create a duplicate intent.")

    add_heading(doc, "5.4  Offline Payments", level=2)
    add_paragraph(
        doc,
        "Not every payment goes through PayMongo. If the tenant hands the landlord cash, or "
        "transfers via direct bank, the landlord uses the offline recording flow:",
    )
    add_bullet(doc, "Landlord enters amount, method (cash / bank / gcash / paymaya / grab_pay / other), optional billing month, and a note.")
    add_bullet(doc, "App calls landlord-record-offline-payment.")
    add_bullet(doc, "Edge Function verifies the caller is the landlord on the contract.")
    add_bullet(doc, "Inserts paired rows into payment and payment_transactions. The paymongo_payment_intent_id is set to 'offline_<uuid>' so it doesn't collide with real PayMongo records.")
    add_bullet(doc, "recorded_by is set to the landlord's id — this is how the ledger distinguishes self-paid online from landlord-recorded offline.")

    add_heading(doc, "5.5  Mock Payments", level=2)
    add_paragraph(
        doc,
        "For demos and testing, the entire PayMongo round-trip can be skipped. This is gated "
        "by a flag called MOCK_PAYMENTS_ENABLED, set via --dart-define at build time.",
    )
    add_bullet(doc, "When enabled, the checkout screen shows a \"Use mock payment\" toggle.")
    add_bullet(doc, "The tenant picks a saved payment_method that has is_mock=true.")
    add_bullet(doc, "App calls paymongo-record-mock-payment.")
    add_bullet(doc, "Edge Function writes synthetic rows directly to payment and payment_transactions, with paymongo_payment_intent_id='pi_mock_<uuid>'.")
    add_bullet(doc, "Move-in flow still flips contract and listing status, so demos look real end-to-end.")

    add_heading(doc, "5.6  The Webhook Safety Net", level=2)
    add_paragraph(
        doc,
        "Webhooks are server-to-server callbacks. PayMongo sends events to the public "
        "endpoint paymongo-webhook whenever a payment succeeds, fails, or is refunded. The "
        "webhook is a critical safety net because mobile networks can drop the connection at "
        "the worst possible moment — between PayMongo confirming the charge and the app "
        "calling paymongo-record-payment.",
    )
    add_paragraph(doc, "Key safety properties of the webhook handler:", bold=True)
    add_kv_bullet(doc, "Idempotent:", "Replays of the same event are safe; nothing gets written twice.")
    add_kv_bullet(doc, "Never downgrades:", "A 'succeeded' row cannot be moved back to 'pending' or 'failed' by a delayed event.")
    add_kv_bullet(doc, "Signature verified:", "Uses PAYMONGO_WEBHOOK_SECRET to confirm events come from PayMongo, not an attacker.")
    add_kv_bullet(doc, "Handles links:", "link.payment.paid events trigger the same payment + payment_transactions writes and mark the payment_links row as 'paid'.")

    add_heading(doc, "5.7  Dual-write to payment and payment_transactions", level=2)
    add_paragraph(
        doc,
        "VXR maintains two payment tables. The original payment table dates from an earlier "
        "Stripe-based implementation; the newer payment_transactions table is the proper "
        "per-attempt ledger used going forward. Until the rest of the app is migrated, every "
        "write goes to both:",
    )
    add_bullet(doc, "payment is the legacy view — joined into older queries and the contract_rent_status view.")
    add_bullet(doc, "payment_transactions is the audit trail — preserves every attempt (including pending and failed), supports offline payments, and is the canonical record for new reports.")
    add_bullet(doc, "Both tables share paymongo_payment_intent_id, so they can be reconciled.")

    add_page_break(doc)


def build_part_6_interactions(doc) -> None:
    add_heading(doc, "Part 6 — Interaction Diagrams", level=1)
    add_paragraph(
        doc,
        "These tables describe the major flows as sequence diagrams. Each row is one step; "
        "the columns show which side (Tenant, Backend, or Landlord) acts in that step.",
    )

    add_heading(doc, "6.1  Tenant Onboarding → First Booking", level=2)
    add_table(
        doc,
        headers=["#", "Tenant", "Backend", "Landlord"],
        rows=[
            ["1", "Signs up with email + password", "Creates auth.users row, ensures profiles row (role=tenant, is_verified=false)", "—"],
            ["2", "Uploads ID + selfie for verification", "verify-identity Edge Function checks faces; writes verifications + flips profiles.is_verified", "—"],
            ["3", "Browses listings", "SELECT on listings + listing_image", "—"],
            ["4", "Taps Apply on a listing", "Verification gate checked", "—"],
            ["5", "Fills 5-step application form", "INSERT into application (status=pending) + application_documents", "Notification trigger fires"],
            ["6", "—", "—", "Reviews application, taps Approve"],
            ["7", "—", "application.status → approved; contract row created (status=awaiting_tenant); tenant notified", "—"],
            ["8", "Reviews and signs contract", "contract.tenant_signature saved; status → awaiting_landlord", "Notification trigger fires"],
            ["9", "—", "—", "Signs contract"],
            ["10", "—", "contract.status → fully_signed", "—"],
            ["11", "Pays move-in (rent + deposit + advance)", "paymongo-create-payment-intent → paymongo-record-payment; contract.status → paid; listings.status → rented", "Notification: payment received"],
            ["12", "Lands on In-Stay Dashboard", "Active contract loaded", "Tenant appears in Tenant Management hub"],
        ],
        col_widths=[1.0, 4.5, 7.5, 4.5],
    )

    add_heading(doc, "6.2  Recurring Rent Payment", level=2)
    add_table(
        doc,
        headers=["#", "Tenant", "Backend", "Landlord"],
        rows=[
            ["1", "Opens Payments tab; sees \"Pay May 2026\" tile", "fetchMyNextDueContract reads contract + computes due months", "—"],
            ["2", "Selects saved or new method", "—", "—"],
            ["3", "Confirms payment", "paymongo-create-payment-intent (billing_month set, amount=monthly_rent)", "—"],
            ["4", "Completes redirect / 3DS", "PayMongo charges method", "—"],
            ["5", "App returns from redirect", "paymongo-record-payment writes payment + payment_transactions; contract status unchanged", "Notification: rent received"],
            ["6", "PaymentScreen shows month as Paid", "—", "Tenant Management updates"],
        ],
        col_widths=[1.0, 5.0, 7.5, 4.0],
    )

    add_heading(doc, "6.3  Maintenance Report Cycle", level=2)
    add_table(
        doc,
        headers=["#", "Tenant", "Backend", "Landlord"],
        rows=[
            ["1", "Submits report (type, priority, photos)", "INSERT into report (status=open); photos uploaded", "Notification: new report"],
            ["2", "—", "—", "Opens report, marks In Progress"],
            ["3", "Notification: status updated", "report.status=in_progress", "—"],
            ["4", "—", "—", "Responds and marks Resolved"],
            ["5", "Sees resolution and response", "report.status=resolved; landlord_response saved", "—"],
        ],
        col_widths=[1.0, 5.0, 7.5, 4.0],
    )

    add_heading(doc, "6.4  Offline Payment Recording", level=2)
    add_table(
        doc,
        headers=["#", "Tenant", "Backend", "Landlord"],
        rows=[
            ["1", "Hands landlord cash for May rent", "—", "Receives money"],
            ["2", "—", "—", "Opens tenant card → Record Payment"],
            ["3", "—", "—", "Enters amount, method=cash, billing_month=May 2026"],
            ["4", "—", "landlord-record-offline-payment validates landlord ownership", "—"],
            ["5", "—", "INSERT into payment_transactions (recorded_by=landlord) + payment", "—"],
            ["6", "Sees May 2026 marked as Paid in PaymentScreen", "—", "Tenant Management ledger updated"],
        ],
        col_widths=[1.0, 4.5, 7.0, 5.0],
    )

    add_heading(doc, "6.5  Termination Cycle", level=2)
    add_table(
        doc,
        headers=["#", "Tenant", "Backend", "Landlord"],
        rows=[
            ["1", "Or: requests early termination", "INSERT contract_termination_request; contract.status=terminating", "Or: initiates termination (mutual / non-renewal / eviction)"],
            ["2", "Sees move-out checklist", "Notifications fire on both sides", "Sees pending termination"],
            ["3", "Vacates property by effective_date", "—", "Inspects unit"],
            ["4", "—", "Final rent settlement and deposit refund recorded in payment ledger", "Confirms move-out"],
            ["5", "Contract appears as Ended in history", "contract.status=ended or terminated", "Listing can be re-published"],
        ],
        col_widths=[1.0, 4.5, 7.0, 5.0],
    )

    add_page_break(doc)


def build_part_7_states(doc) -> None:
    add_heading(doc, "Part 7 — State Reference Tables", level=1)
    add_paragraph(
        doc,
        "These are the lifecycle states for each major entity. Knowing them is essential for "
        "answering \"what's the next valid step?\" at any point in the system.",
    )

    add_heading(doc, "7.1  Application Statuses", level=2)
    add_table(
        doc,
        headers=["Status", "Meaning", "Next state(s)"],
        rows=[
            ["pending", "Tenant submitted; landlord hasn't decided yet", "approved, rejected, withdrawn"],
            ["approved", "Landlord accepted; contract row created", "(terminal for application; contract takes over)"],
            ["rejected", "Landlord declined", "(terminal — tenant can re-apply once contract is closed)"],
            ["withdrawn", "Tenant pulled the application before decision", "(terminal)"],
        ],
    )

    add_heading(doc, "7.2  Contract Statuses", level=2)
    add_table(
        doc,
        headers=["Status", "Meaning", "Next state(s)"],
        rows=[
            ["awaiting_tenant", "Contract created from approved application; tenant must sign", "awaiting_landlord, cancelled"],
            ["awaiting_landlord", "Tenant signed; landlord must sign", "fully_signed, cancelled"],
            ["fully_signed", "Both signed; awaiting move-in payment", "paid, cancelled"],
            ["paid", "Move-in paid; tenancy is active", "terminating, expiring, ended, terminated"],
            ["terminating", "One party requested early termination", "ended, terminated"],
            ["expiring", "Lease nearing end of fixed term", "ended"],
            ["ended", "Lease ran its full course", "(terminal)"],
            ["terminated", "Lease ended early", "(terminal)"],
            ["cancelled", "Contract abandoned before signing or payment", "(terminal)"],
        ],
    )

    add_heading(doc, "7.3  Payment Statuses", level=2)
    add_table(
        doc,
        headers=["Status", "Meaning"],
        rows=[
            ["pending", "Intent created, awaiting method attach"],
            ["requires_action", "Awaiting 3DS or e-wallet redirect"],
            ["succeeded", "Money received"],
            ["failed", "PayMongo declined the charge"],
            ["refunded", "Money returned to tenant"],
            ["cancelled", "Intent cancelled before completion"],
        ],
    )

    add_heading(doc, "7.4  Report Statuses", level=2)
    add_table(
        doc,
        headers=["Status", "Meaning", "Who can move it"],
        rows=[
            ["open", "Tenant just filed", "Landlord can move to in_progress; tenant can cancel"],
            ["in_progress", "Landlord acknowledged and is working on it", "Landlord can move to resolved"],
            ["resolved", "Issue fixed", "(terminal)"],
            ["cancelled", "Tenant withdrew or duplicate report", "(terminal)"],
        ],
    )

    add_heading(doc, "7.5  Verification Decisions", level=2)
    add_table(
        doc,
        headers=["Decision", "Meaning", "Effect on profile"],
        rows=[
            ["approved", "AI confirmed identity match", "profiles.is_verified = true"],
            ["rejected", "AI denied (faces don't match, doc unreadable, etc.)", "profiles.is_verified stays false"],
            ["manual_review", "AI is unsure; admin must review", "profiles.is_verified stays false until admin acts"],
        ],
    )

    add_page_break(doc)


def build_part_8_lookup(doc) -> None:
    add_heading(doc, "Part 8 — Quick Lookup: Which Screen Touches What?", level=1)
    add_paragraph(
        doc,
        "Use this table as a cheat-sheet during review. For each screen, the table lists the "
        "Supabase tables it reads or writes, the Edge Functions it calls, and the storage "
        "buckets it touches.",
    )
    add_table(
        doc,
        headers=["Screen", "Reads / Writes", "Edge Functions", "Storage"],
        rows=[
            ["LandingPage / AuthGate", "auth.users", "—", "—"],
            ["SignUpScreen / LoginScreen", "auth.users, profiles", "—", "—"],
            ["VerificationScreen", "verifications, profiles", "verify-identity", "verifications"],
            ["HomeScreen", "listings, listing_image, bookmarks", "—", "listing-images (read)"],
            ["SearchScreen", "listings, listing_image, bookmarks", "—", "listing-images (read)"],
            ["UnitDetailsScreen", "listings, listing_pricing, listing_lease_term, listing_amenities, listing_image, application", "—", "listing-images (read)"],
            ["FavoritesScreen", "bookmarks, listings", "—", "—"],
            ["RentalApplicationScreen", "application, application_documents", "—", "application-docs"],
            ["MyApplicationsScreen", "application", "—", "—"],
            ["ContractViewScreen", "contract, listings, application, profiles", "—", "listing-contracts"],
            ["ContractPaymentScreen", "payment, payment_methods, payment_transactions, contract", "paymongo-create-payment-intent, paymongo-attach-payment-method, paymongo-record-payment, paymongo-record-mock-payment", "—"],
            ["PaymentScreen", "payment, payment_methods, payment_transactions, contract", "paymongo-add-method, paymongo-delete-method, paymongo-set-default-method", "—"],
            ["InStayDashboardScreen", "contract, payment, report, contract_termination_request, conversation", "—", "—"],
            ["ReportManagementScreen", "report", "—", "Report photo storage"],
            ["ConversationsScreen", "conversation, profiles, message", "—", "—"],
            ["ChatThreadScreen", "message, conversation", "—", "conversation-attachments"],
            ["HouseEnlistmentScreen / CreateListing", "listings, listing_pricing, listing_lease_term, listing_amenities", "—", "verifications (landlord docs)"],
            ["ListingImageManagerScreen", "listing_image, listings (cover photo)", "—", "listing-images"],
            ["PanoramaCaptureScreen / Manager", "listing_image", "—", "listing-images"],
            ["ApplicantsScreen / ApplicationsScreen", "application, listings, application_documents", "—", "—"],
            ["LandlordContractsScreen", "contract, listings, application", "—", "—"],
            ["TenantManagementScreen", "contract, application, listings, payment_transactions, payment_links, report", "landlord-record-offline-payment, paymongo-create-payment-link", "—"],
        ],
        col_widths=[4.5, 7.0, 4.5, 2.5],
    )

    add_page_break(doc)


def build_part_9_glossary(doc) -> None:
    add_heading(doc, "Part 9 — Glossary", level=1)
    add_paragraph(
        doc,
        "Plain-language definitions of the technical terms used throughout this reviewer.",
    )

    add_kv_bullet(doc, "Auth (Authentication):", "The process of proving who a user is. In VXR, this means signing in with email + password (or Google) and getting back a JWT.")
    add_kv_bullet(doc, "JWT (JSON Web Token):", "A small encoded blob the app sends with every request. It tells Supabase \"I am user X and I logged in N minutes ago.\" Used to enforce RLS.")
    add_kv_bullet(doc, "RLS (Row Level Security):", "A Postgres feature that filters query results based on rules tied to the current user. Even if the app asks for someone else's row, RLS hides it.")
    add_kv_bullet(doc, "Edge Function:", "A small TypeScript program that runs on Supabase's servers. Used for tasks that need a secret API key or must succeed/fail atomically.")
    add_kv_bullet(doc, "Realtime:", "Live websocket updates from the database. When a row is inserted, all subscribed clients see it within milliseconds.")
    add_kv_bullet(doc, "Webhook:", "A URL that an external service calls to notify your backend that something happened (e.g. PayMongo notifying VXR that a payment succeeded).")
    add_kv_bullet(doc, "Idempotency:", "A property where the same request can be replayed safely without producing duplicate effects. The PayMongo webhook is idempotent — replays don't double-charge or double-record.")
    add_kv_bullet(doc, "Storage Bucket:", "A folder-like container in Supabase Storage. Each bucket has its own access rules.")
    add_kv_bullet(doc, "PayMongo Intent vs Link:", "An Intent is a programmatic charge handled inside the app. A Link is a hosted checkout URL the user opens in a browser. VXR uses Intents for in-app payments and Links for landlord-generated rent invoices.")
    add_kv_bullet(doc, "Mock Payment:", "A fake payment that skips PayMongo. Used in demos to look real without spending money. Enabled by the MOCK_PAYMENTS_ENABLED flag.")
    add_kv_bullet(doc, "Dual-write:", "Writing the same data to two tables at once for backwards compatibility. VXR dual-writes to payment (legacy) and payment_transactions (new ledger).")
    add_kv_bullet(doc, "PSGC (Philippine Standard Geographic Code):", "Official codes for Philippine regions, provinces, cities, and barangays. VXR uses PSGC JSON files to power the location picker.")
    add_kv_bullet(doc, "Haversine formula:", "A math formula that calculates the great-circle distance between two latitude/longitude points on a sphere. VXR uses it to show \"distance to property\" on listing cards without calling an API.")
    add_kv_bullet(doc, "Sandbox vs Production:", "PayMongo has two sets of API keys. Sandbox uses fake money for testing; production uses real money. VXR is currently sandbox.")
    add_kv_bullet(doc, "Billing month:", "A YYYY-MM-01 date that identifies which monthly rent a payment covers. Used to prevent double-payment of the same month.")
    add_kv_bullet(doc, "Cover photo:", "The single image from a listing that appears on cards in the home feed and search results.")

    add_paragraph(doc, "")
    add_paragraph(doc, "End of reviewer document.", italic=True)


# ---------------------------------------------------------------------------
# Build & save
# ---------------------------------------------------------------------------

def configure_document(doc) -> None:
    # Page margins
    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # Default font on Normal style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def main() -> None:
    doc = Document()
    configure_document(doc)

    build_cover(doc)
    build_toc_page(doc)
    build_part_1_overview(doc)
    build_part_2_tenant(doc)
    build_part_3_landlord(doc)
    build_part_4_backend(doc)
    build_part_5_payments(doc)
    build_part_6_interactions(doc)
    build_part_7_states(doc)
    build_part_8_lookup(doc)
    build_part_9_glossary(doc)

    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    out_path = os.path.join(project_root, "VXR_System_Flow_Reviewer.docx")
    doc.save(out_path)

    size_kb = os.path.getsize(out_path) / 1024.0
    print(f"Wrote: {out_path}")
    print(f"Size:  {size_kb:.1f} KB")


if __name__ == "__main__":
    main()
